from tkinter import *
from tkinter import ttk
from tkinter import filedialog
import os
import pandas as pd
from docx import Document
from openpyxl import load_workbook

doc = Document()
arquivo = load_workbook('OpenPy.xlsx')
planilha = arquivo.active



# setando valor do cabeçalho
cabecalho = "teste numero 2349801823094814"

def adicionar_cabecalho(variavel):
    doc.add_heading(variavel, level=1)
    doc.save("teste.docx")
    print(f"Valor atual da variável 'cabeçalho': {cabecalho}")
    print("Relatório salvo com sucesso!")



def pesquisa():
    print("BOTÃO DE PESQUISA CLICADO")


def open_new_window():
    selected_item = tree.focus()  # Get selected item ID
    if selected_item:
        values = tree.item(selected_item, 'values')  # Get row data

        # Create new window
        new_win = Toplevel(root)
        new_win.title("Detalhes")

        # Create a frame for the canvas and scrollbar
        frame = Frame(new_win)
        frame.pack(fill=BOTH, expand=True)

        # Create a canvas inside the frame
        canvas = Canvas(frame)
        canvas.pack(side=LEFT, fill=BOTH, expand=True)

        # Add a scrollbar to the canvas
        scrollbar = ttk.Scrollbar(frame, orient=VERTICAL, command=canvas.yview)
        scrollbar.pack(side=RIGHT, fill=Y)



        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1*(event.delta/120)), "units")

        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.bind_all("<MouseWheel>", _on_mousewheel)



        # Create another frame inside the canvas for the content
        content_frame = Frame(canvas)
        canvas.create_window((0, 0), window=content_frame, anchor='nw')

        column_names = tree["columns"]

        # Add the labels to the content frame
        for i, val in enumerate(values):
            column_name = column_names[i]
            
            # Column name label
            Label(content_frame, text=column_name, font=('Arial', 10, 'bold')).grid(row=i, column=0, sticky='w', padx=10, pady=5)
            
            # Selectable value using readonly Entry
            entry = ttk.Entry(content_frame, font=('Arial', 10), width=30)
            entry.insert(0, val)
            entry.config(state='readonly')
            entry.grid(row=i, column=1, sticky='w', padx=10, pady=5)


            Button(content_frame, 
                   text="Adicionar ao relatório", 
                   command=lambda cn=column_name, 
                   v=val: print(f"{cn}: {v}")
                   
                   ).grid(row=i, column=2, padx=10, pady=5)

            Button(content_frame, 
                   text="Cabeçalho", 
                   command=lambda v=val: adicionar_cabecalho(v)
                   ).grid(row=i, column=3, padx=10, pady=5)
             
        # Update scrollregion when contents are added
        content_frame.update_idletasks()
        canvas.config(scrollregion=canvas.bbox("all"))

    else:
        print("Nenhuma linha selecionada.")


# carregar arquivo
def load_file():
    caminho_arquivo = filedialog.askopenfilename()

    # Caso o arquivo seja carregado, verificar se é um arquivo de planilha
    if caminho_arquivo:
        nome_arquivo, extensao_arquivo = os.path.splitext(caminho_arquivo)

        if extensao_arquivo == '.csv':
            try:
                df = pd.read_csv(caminho_arquivo, encoding='utf-8', sep=';')
                print("Arquivo de planilha carregado")

            except UnicodeDecodeError:
                df = pd.read_csv(caminho_arquivo, encoding='ISO-8859-1', sep=';')
                print("Arquivo de planilha carregado")

        elif extensao_arquivo in ['.xlsx', '.xls']:
            df = pd.read_excel(caminho_arquivo)
            print("Arquivo de planilha carregado")

        # filter data from excel using pandas


        else:
            print("tipo de arquivo não permitido")
            return

        # limpar dados do treeview
        for col in tree.get_children():
            tree.delete(col)
        
        # definir colunas do treeview
        
        tree["columns"] = list(df.columns)
        tree["show"] = "headings" # oculta a primeira coluna vazia

        # configura os títulos das colunas
        for coluna in df.columns:
            tree.heading (coluna, text=coluna)
            tree.column(coluna, anchor='center', minwidth=100)

        
        # insere os dados no Treeview
        for index, row in df.iterrows():
            tree.insert("", "end", values=list(row))


# centralizar window
def center_window(width=300, height=200):

    # get screen width and height
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()

    # calculate position x and y coordinates
    x = (screen_width/2) - (width/2)
    y = (screen_height/2) - (height/2)
    root.geometry('%dx%d+%d+%d' % (width, height, x, y))


root = Tk()
root.title("Planilha")
center_window(500, 400)
root.resizable(True, True)

root.grid_rowconfigure(2, weight=1)
root.grid_columnconfigure(0, weight=1)

# Use grid for all root widgets
entrada_texto = ttk.Entry(root, width=20) # =======================================================================
entrada_texto.grid(row=0, column=0, padx=5, pady=10)

pesquisa_btn = ttk.Button(root, text="Pesquisar", command=pesquisa)
pesquisa_btn.grid(row=0, column=1, padx=1, pady=1)

detalhes_btn = ttk.Button(root, text="Detalhes", command=open_new_window)
detalhes_btn.grid(row=0, column=2, padx=5, pady=10)

carregar_btn = ttk.Button(root, text="Carregar Arquivo", command=load_file)
carregar_btn.grid(row=0, column=3, padx=5, pady=10)

# Use pack/grid for frame as before
frame = ttk.Frame(root)
frame.grid(row=2, column=0, columnspan=3, sticky="nsew", padx=10, pady=10)

# Scrollbars
scroll_y = ttk.Scrollbar(frame, orient="vertical")
scroll_x = ttk.Scrollbar(frame, orient="horizontal")

tree = ttk.Treeview(frame, yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)

# Use grid instead of place
frame.columnconfigure(0, weight=1)
frame.rowconfigure(0, weight=1)

tree.grid(row=0, column=0, sticky="nsew")
scroll_y.grid(row=0, column=1, sticky="ns")
scroll_x.grid(row=1, column=0, sticky="ew")

scroll_y.config(command=tree.yview)
scroll_x.config(command=tree.xview)

root.mainloop()